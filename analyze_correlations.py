"""
analyze_correlations.py
=========================
Merges Career_Totals (win/loss records) with Bot_Features (weapon type,
drivetrain, weight, etc.) and computes, for every feature value, how much
better or worse robots with that feature performed compared to the overall
average. This is the "does X feature show up more often among winners"
question, made numeric.

METHOD (deliberately simple, and worth explaining this way in your report):

  For a feature like Weapon_Type, and a value like "Horizontal Spinner":
    group_win_rate   = average Career_Win% of robots whose Weapon_Type
                        contains "Horizontal Spinner"
    overall_win_rate = average Career_Win% across ALL robots in the dataset
    lift             = group_win_rate / overall_win_rate

  lift > 1.0  -> robots with that feature won MORE than the field average
  lift < 1.0  -> robots with that feature won LESS than the field average
  lift = 1.0  -> no difference

This is a simple ratio-of-averages ("lift"), not a formal statistical test —
it does NOT prove causation, and small sample sizes (e.g. only 3 robots ever
used a given weapon type) will produce noisy, unreliable lift scores. The
Robot_Count column is included specifically so you can filter those out
before drawing conclusions in your report (a common rule of thumb: don't
trust a group with fewer than ~5-10 robots).

Usage:
    python analyze_correlations.py
    (run battlebots_scraper.py AND bot_features_scraper.py first)

Output:
    Adds a "Feature_Correlations" sheet to battlebots_2015_2025.xlsx
    Also prints the top 10 highest- and lowest-lift features to the console.
"""

import re
import pandas as pd
from openpyxl import load_workbook

WORKBOOK = "battlebots_2015_2025.xlsx"

# Which columns from Bot_Features to test for a win-rate relationship.
# Categorical / short-text columns work best; free-text ones will just
# produce lots of tiny one-robot groups (still shown, but low Robot_Count).
FEATURE_COLUMNS = ["Weapon_Type", "Drive_Type", "Wheel_Count", "Weight_Class"]

MIN_GROUP_SIZE_WARNING = 5  # flagged in output, not filtered out


def load_merged():
    career = pd.read_excel(WORKBOOK, sheet_name="Career_Totals")
    features = pd.read_excel(WORKBOOK, sheet_name="Bot_Features")
    merged = career.merge(features, on="Robot", how="inner")
    return merged


def extract_values(cell, column):
    """
    Turn one cell's text into a list of comparable "tags".
    Most feature text (e.g. 'Horizontal Bar Spinner') is used as one tag.
    Wheel_Count is numeric already, so it's returned as-is.
    """
    if pd.isna(cell):
        return []
    if column == "Wheel_Count":
        return [str(int(cell)) + "-wheel"] if not pd.isna(cell) else []
    text = str(cell)
    # Some infoboxes list multiple weapons/drive notes separated by slashes,
    # commas, or "and" — split those into separate tags so each is counted.
    parts = re.split(r"\s*(?:/|,| and )\s*", text)
    return [p.strip() for p in parts if p.strip()]


def compute_feature_correlations(merged: pd.DataFrame) -> pd.DataFrame:
    overall_avg = merged["Career_Win %"].mean()
    rows = []

    for col in FEATURE_COLUMNS:
        if col not in merged.columns:
            continue

        # Build a long-format (tag, win%) table so multi-value cells all count
        tagged = []
        for _, r in merged.iterrows():
            for tag in extract_values(r[col], col):
                tagged.append((tag, r["Career_Win %"]))

        if not tagged:
            continue

        tag_df = pd.DataFrame(tagged, columns=["Value", "Career_Win %"])
        grouped = tag_df.groupby("Value").agg(
            Robot_Count=("Career_Win %", "count"),
            Avg_Win_Pct=("Career_Win %", "mean"),
        ).reset_index()

        grouped["Feature"] = col
        grouped["Overall_Avg_Win_Pct"] = round(overall_avg, 3)
        grouped["Lift"] = (grouped["Avg_Win_Pct"] / overall_avg).round(3)
        grouped["Avg_Win_Pct"] = grouped["Avg_Win_Pct"].round(3)
        grouped["Small_Sample_Warning"] = grouped["Robot_Count"] < MIN_GROUP_SIZE_WARNING

        rows.append(grouped[["Feature", "Value", "Robot_Count", "Avg_Win_Pct",
                              "Overall_Avg_Win_Pct", "Lift", "Small_Sample_Warning"]])

    if not rows:
        return pd.DataFrame()

    result = pd.concat(rows, ignore_index=True)
    result = result.sort_values(["Feature", "Lift"], ascending=[True, False]).reset_index(drop=True)
    return result


def append_sheet(df, workbook=WORKBOOK, sheet_name="Feature_Correlations"):
    wb = load_workbook(workbook)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
        wb.save(workbook)
    with pd.ExcelWriter(workbook, engine="openpyxl", mode="a") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
    print(f"'{sheet_name}' sheet written to {workbook}.")


if __name__ == "__main__":
    merged = load_merged()
    result = compute_feature_correlations(merged)

    if result.empty:
        print("No overlapping data between Career_Totals and Bot_Features. "
              "Did bot_features_scraper.py run successfully?")
    else:
        append_sheet(result)

        reliable = result[~result["Small_Sample_Warning"]]
        print("\nTop 10 features associated with HIGHER win rate (Robot_Count >= 5):")
        print(reliable.sort_values("Lift", ascending=False).head(10)
              [["Feature", "Value", "Robot_Count", "Avg_Win_Pct", "Lift"]]
              .to_string(index=False))

        print("\nTop 10 features associated with LOWER win rate (Robot_Count >= 5):")
        print(reliable.sort_values("Lift", ascending=True).head(10)
              [["Feature", "Value", "Robot_Count", "Avg_Win_Pct", "Lift"]]
              .to_string(index=False))
