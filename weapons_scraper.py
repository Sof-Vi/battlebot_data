"""
battlebots_com_weapons_scraper.py
===================================
Fills the weapon-type gap for seasons 8-12 (2018-2022). Wikipedia's
Contestants table only has a Weapon column for seasons 6-7 (2015-2016) --
from season 8 onward it only has Fight Record, no weapon info at all. The
Fandom wiki infobox (bot_features_scraper.py) has weapon type but coverage
is inconsistent robot-to-robot.

battlebots.com itself has a "season robots" page per year, and every robot's
card on that page includes an official "Type:" field (its weapon), e.g.:

    Robot: Big Dill  Builder: Emmanuel Carrillo  Type: Lifter  Job: ...

This script fetches those pages directly from the source and uses them as
the authoritative weapon type, since BattleBots Inc. itself is the one
publishing it (more complete and more "official" than either Wikipedia or
the fan wiki for this specific field).

Usage:
    python battlebots_com_weapons_scraper.py
    (run battlebots_scraper.py and bot_features_scraper.py first)

Output:
    - Adds a "Weapons_BattleBotsCom" sheet (every season's weapon listing, raw)
    - Updates the "Bot_Features" sheet:
        Weapon_Type              -> now prefers battlebots.com, falls back to
                                     whatever bot_features_scraper.py found on
                                     Fandom if battlebots.com has nothing
        Weapon_Type_Fandom       -> the original Fandom-only value, kept for
                                     comparison/audit
        Weapon_Type_BattleBotsCom-> the official value, most recent season
        Weapon_History_BattleBotsCom -> every distinct weapon type across all
                                     seasons that robot competed in (useful if
                                     a robot was rebuilt with a new weapon)
"""

import os
import re
import sys
import time

import requests
import pandas as pd
from bs4 import BeautifulSoup

from battlebots_scraper import robot_key  # reuse the same name-matching logic

WORKBOOK = "battlebots_2015_2025.xlsx"
HEADERS = {
    "User-Agent": "BattleBotsStatsResearch/1.0 (student research project; contact: youremail@example.com)"
}
MAX_RETRIES = 3

# (page URL, season label -- MUST match battlebots_scraper.py's SEASONS labels, year)
SEASON_PAGES = [
    ("https://battlebots.com/season-1-robots/",            "Season 6 (2015)",  2015),
    ("https://battlebots.com/season-2-robots/",             "Season 7 (2016)",  2016),
    ("https://battlebots.com/2018-season-robots/",          "Season 8 (2018)",  2018),
    ("https://battlebots.com/2019-season-robots/",          "Season 9 (2019)",  2019),
    ("https://battlebots.com/2020-season-robots/",          "Season 10 (2020)", 2020),
    ("https://battlebots.com/2021-season-robots/",          "Season 11 (2021)", 2021),
    ("https://battlebots.com/world-championship-vii-robots/","Season 12 (2022)", 2022),
    # If BattleBots.com publishes a page for a newer season, add it here, e.g.:
    # ("https://battlebots.com/world-championship-viii-robots/", "Season 13 (2024)", 2024),
]

# Matches "Robot: <name> Builder: <name> Type: <weapon> Job:" in the page's
# flattened text. Fields are checked in this fixed order because that's the
# order battlebots.com's own page template lists them in.
ROBOT_BLOCK_RE = re.compile(
    r"Robot:\s*(?P<robot>.+?)\s+Builder:\s*(?P<builder>.+?)\s+Type:\s*(?P<weapon>.+?)\s+Job:",
    re.IGNORECASE,
)


def fetch_page_text(url: str) -> str:
    """Fetch a battlebots.com page and flatten it to plain text (strip nav/
    footer/scripts first so they can't produce false regex matches)."""
    last_error = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.get(url, headers=HEADERS, timeout=30)
        except requests.exceptions.RequestException as e:
            last_error = f"network error ({e.__class__.__name__}): {e}"
            print(f"    attempt {attempt}/{MAX_RETRIES} failed: {last_error}")
            time.sleep(1.5 * attempt)
            continue
        if resp.status_code != 200:
            last_error = f"HTTP {resp.status_code}"
            print(f"    attempt {attempt}/{MAX_RETRIES} failed: {last_error}")
            time.sleep(1.5 * attempt)
            continue

        soup = BeautifulSoup(resp.text, "lxml")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        return soup.get_text(separator=" ", strip=True)

    raise ConnectionError(f"Could not fetch '{url}' after {MAX_RETRIES} attempts. Last error: {last_error}")


def parse_season_page(text: str, label: str):
    rows = []
    for m in ROBOT_BLOCK_RE.finditer(text):
        robot = m.group("robot").strip()
        weapon = m.group("weapon").strip()
        if not robot or not weapon:
            continue
        rows.append({"Robot": robot, "Weapon_Type": weapon})
    if not rows:
        print(f"    WARNING: found zero robot cards on '{label}'. The page's "
              f"template may have changed -- check the URL manually.")
    return rows


def build_weapons_log():
    all_rows = []
    for url, label, year in SEASON_PAGES:
        print(f"Fetching {label} from {url} ...")
        try:
            text = fetch_page_text(url)
            rows = parse_season_page(text, label)
        except Exception as e:
            print(f"  FAILED: {label} -> {e}\n")
            continue
        for r in rows:
            r["Season"] = label
            r["Year"] = year
        all_rows.extend(rows)
        print(f"  OK: {label} -> {len(rows)} robots")
        time.sleep(0.5)

    if not all_rows:
        sys.exit(
            "\nERROR: No weapon data collected from battlebots.com at all.\n"
            "Check the FAILED/WARNING lines above -- likely a network issue, "
            "or battlebots.com restructured its season-robots page template.\n"
        )

    df = pd.DataFrame(all_rows)[["Season", "Year", "Robot", "Weapon_Type"]]
    df["Robot_Key"] = df["Robot"].map(robot_key)
    return df


def merge_into_bot_features(weapons_log: pd.DataFrame, workbook=WORKBOOK):
    if not os.path.exists(workbook):
        sys.exit(f"\nERROR: '{workbook}' not found. Run battlebots_scraper.py first.\n")
    try:
        features = pd.read_excel(workbook, sheet_name="Bot_Features")
    except ValueError:
        sys.exit(f"\nERROR: no 'Bot_Features' sheet in '{workbook}'. Run bot_features_scraper.py first.\n")

    features["Robot_Key"] = features["Robot"].map(robot_key)

    # Most recent season's weapon = this robot's "current" official weapon
    latest = (
        weapons_log.sort_values("Year")
        .groupby("Robot_Key", as_index=False)
        .last()[["Robot_Key", "Weapon_Type"]]
        .rename(columns={"Weapon_Type": "Weapon_Type_BattleBotsCom"})
    )
    # Every distinct weapon type across all seasons (shows rebuilds/weapon swaps)
    history = (
        weapons_log.groupby("Robot_Key")["Weapon_Type"]
        .apply(lambda s: " / ".join(sorted(set(s))))
        .reset_index()
        .rename(columns={"Weapon_Type": "Weapon_History_BattleBotsCom"})
    )

    if "Weapon_Type" in features.columns:
        features = features.rename(columns={"Weapon_Type": "Weapon_Type_Fandom"})
    else:
        features["Weapon_Type_Fandom"] = None

    features = features.merge(latest, on="Robot_Key", how="left")
    features = features.merge(history, on="Robot_Key", how="left")

    # Final Weapon_Type: prefer the official battlebots.com value; fall back
    # to whatever Fandom had if battlebots.com didn't cover that robot.
    features["Weapon_Type"] = features["Weapon_Type_BattleBotsCom"].combine_first(
        features["Weapon_Type_Fandom"]
    )

    filled = features["Weapon_Type_BattleBotsCom"].notna().sum()
    total = len(features)
    print(f"\nMatched {filled}/{total} robots in Bot_Features to an official "
          f"battlebots.com weapon type.")
    unmatched = features.loc[features["Weapon_Type_BattleBotsCom"].isna(), "Robot"].tolist()
    if unmatched:
        print(f"  {len(unmatched)} robot(s) with no battlebots.com match "
              f"(name spelling differences, or robot not on a scraped season page):")
        print(f"  {unmatched[:15]}{' ...' if len(unmatched) > 15 else ''}")

    front = ["Robot", "Weapon_Type", "Weapon_Type_BattleBotsCom", "Weapon_Type_Fandom",
             "Weapon_History_BattleBotsCom"]
    rest = [c for c in features.columns if c not in front + ["Robot_Key"]]
    features = features[[c for c in front if c in features.columns] + rest]
    return features


def write_results(weapons_log, features, workbook=WORKBOOK):
    from openpyxl import load_workbook

    wb = load_workbook(workbook)
    for sheet in ("Weapons_BattleBotsCom", "Bot_Features"):
        if sheet in wb.sheetnames:
            del wb[sheet]
    wb.save(workbook)

    with pd.ExcelWriter(workbook, engine="openpyxl", mode="a") as writer:
        weapons_log.drop(columns=["Robot_Key"]).to_excel(
            writer, sheet_name="Weapons_BattleBotsCom", index=False)
        features.to_excel(writer, sheet_name="Bot_Features", index=False)

    import excel_style
    excel_style.style_workbook(workbook)
    print(f"\n'Weapons_BattleBotsCom' and updated 'Bot_Features' written to {workbook}.")


if __name__ == "__main__":
    weapons_log = build_weapons_log()
    features = merge_into_bot_features(weapons_log)
    write_results(weapons_log, features)